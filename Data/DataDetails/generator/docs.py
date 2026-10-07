"""RFDS (xlsx/pdf), MA/SA pdf, catalog, rules, BOM workbooks, Sitetracker CSVs."""
import hashlib, csv, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
import model as M

SN = M.SECT_NAMES
def h(x, n):  # deterministic small ints
    return int(hashlib.md5(x.encode()).hexdigest(), 16) % n

HDR = PatternFill("solid", fgColor="1F3864"); SUB = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="A6A6A6"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

def _hdr(ws, row, vals, col=1, fill=HDR):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = Font(bold=True, color="FFFFFF" if fill is HDR else "000000", size=9)
        c.fill = fill; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

def _row(ws, row, vals, col=1):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=col + i, value=v); c.border = BOX; c.font = Font(size=9)

def _widths(ws, w):
    for i, x in enumerate(w): ws.column_dimensions[get_column_letter(i + 1)].width = x

def tilts(s, d):
    k = s["site_id"] + d["sector"] + str(d["pos"])
    return h(k + "m", 3), h(k + "l", 6) + 2, h(k + "h", 6) + 1

def dates(s):
    f = dt.date.fromisoformat(s["flight"])
    return dict(rfds_r1=f - dt.timedelta(days=62), rfds_r2=f - dt.timedelta(days=40), rfds_r3=f - dt.timedelta(days=12),
                cd=f + dt.timedelta(days=9), ma=f + dt.timedelta(days=16), flight=f,
                assigned=dt.date(2026, 9, 14) + dt.timedelta(days=h(s["site_id"], 4)), rev0=f + dt.timedelta(days=24))

def rfds_positions(s, state):
    """state 'existing' or 'final' -> list of (sector, pos, antenna_dev, [radio_devs])"""
    out = []
    ok = (lambda d: d["status"] in ("Existing", "Remove")) if state == "existing" else (lambda d: d["status"] in ("Existing", "New"))
    devs = [d for d in s["devices"] if d["in_rfds"] and ok(d)]
    for sc in s["sectors"]:
        for p in (1, 2, 3):
            here = [d for d in devs if d["sector"] == sc["name"] and d["pos"] == p]
            ants = [d for d in here if d["kind"] in ("passive", "air")]
            if not ants: continue
            out.append((sc, p, ants[0], [d for d in here if d["kind"] == "radio"]))
    return out

def site_info(s):
    D = dates(s)
    return [("Site ID", s["site_id"]), ("Site Name", s["name"]), ("Market", s["market"]), ("Customer", s["customer"]),
            ("Address", s["addr"]), ("City", s["city"]), ("State", s["state"]), ("ZIP", s["zip"]),
            ("Latitude", s["lat"]), ("Longitude", s["lon"]), ("Structure Type", s["structure"]),
            ("Structure Height (ft)", s["height"]), ("Structure Owner", s["owner"]), ("Project Type", s["project"]),
            ("RFDS Revision", s["rfds_rev"]), ("RFDS Date", D["rfds_r3"].isoformat()),
            ("RF Engineer", ["J. Alvarez", "M. Chen", "S. Patel", "R. Okafor"][h(s["site_id"], 4)] + " (synthetic)")]

def action(d):
    return {"Existing": "Retain", "New": "Add", "Remove": "Remove"}[d["status"]]

def write_rfds_v2(s, path):
    wb = Workbook(); ws = wb.active; ws.title = "Site Info"
    ws["A1"] = f"RF DATA SHEET - {s['customer']}"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "SYNTHETIC SAMPLE - NOT FOR CONSTRUCTION"; ws["A2"].font = Font(italic=True, color="C00000")
    for i, (k, v) in enumerate(site_info(s)):
        ws.cell(row=4 + i, column=1, value=k).font = Font(bold=True, size=9); ws.cell(row=4 + i, column=2, value=v).font = Font(size=9)
    _widths(ws, [24, 40])
    for state, title in (("existing", "Existing Config"), ("final", "Final Config")):
        w = wb.create_sheet(title)
        w["A1"] = f"{s['site_id']} - {title.upper()}"; w["A1"].font = Font(bold=True, size=12)
        if state == "existing":
            cols = ["Sector", "Azimuth (deg)", "Ant Position", "Antenna Mfr", "Antenna Model", "Rad Center (ft)", "Mech Tilt (deg)",
                    "Radios", "Radio Bands", "Technology"]
        else:
            cols = ["Sector", "Azimuth (deg)", "Ant Position", "Action", "Antenna / AAU Mfr", "Antenna / AAU Model", "Rad Center (ft)",
                    "Mech Tilt (deg)", "E-Tilt LB", "E-Tilt MB/HB", "Radio 1", "Radio 1 Action", "Radio 2", "Radio 2 Action",
                    "Bands", "Technology", "MIMO"]
        _hdr(w, 3, cols); r = 4
        for sc, p, a, rads in rfds_positions(s, state):
            c = M.CATALOG[a["key"]]
            mt, el, eh = tilts(s, a)
            if state == "existing":
                vals = [SN[sc["name"]], sc["az"], p, c["mfr"], c["model"], a["rc"], mt,
                        ", ".join(M.CATALOG[x["key"]]["model"] for x in rads), ", ".join(M.CATALOG[x["key"]]["bands"] for x in rads),
                        "LTE"]
            else:
                if a["kind"] == "air":
                    r1, r1a, r2, r2a, bands, mimo = "Integrated", action(a), "", "", c["bands"], "64T64R"
                else:
                    rr = rads + [None, None]
                    r1 = M.CATALOG[rr[0]["key"]]["model"] if rr[0] else ""; r1a = action(rr[0]) if rr[0] else ""
                    r2 = M.CATALOG[rr[1]["key"]]["model"] if rr[1] else ""; r2a = action(rr[1]) if rr[1] else ""
                    bands = "/".join(M.CATALOG[x["key"]]["bands"] for x in rads); mimo = "4T4R"
                act = action(a)
                if act == "Add" and any(d["sector"] == sc["name"] and d["pos"] == p and d["status"] == "Remove" and d["kind"] == "passive"
                                        for d in s["devices"] if d["in_rfds"]):
                    act = "Swap"
                vals = [SN[sc["name"]], sc["az"], p, act, c["mfr"], c["model"], a["rc"], mt, el, eh, r1, r1a, r2, r2a, bands, c["tech"], mimo]
            _row(w, r, vals); r += 1
        _widths(w, [9, 9, 8] + [14] * (len(cols) - 3))
    w = wb.create_sheet("Baseband & Power")
    _hdr(w, 1, ["Item", "Model", "Action", "Qty", "Notes"]); r = 2
    for bb in s["basebands"]:
        _row(w, r, ["Baseband", M.CATALOG[bb]["model"], "Add", 1, "Rack in existing cabinet"]); r += 1
    _row(w, r, ["Router", "Router 6675", "Retain", 1, ""]); r += 1
    load = sum(M.CATALOG[d["key"]]["power_w"] for d in s["devices"] if d["in_rfds"] and d["status"] == "New")
    _row(w, r, ["DC load added (W)", "", "", load, "Peak, new radios and AAUs"]); _widths(w, [22, 18, 10, 8, 30])
    w = wb.create_sheet("Revision History")
    D = dates(s); _hdr(w, 1, ["Rev", "Date", "Description"])
    for i, (rv, dd, ds) in enumerate([("R1", D["rfds_r1"], "Preliminary"), ("R2", D["rfds_r2"], "Band plan update"), ("R3", D["rfds_r3"], "Final for construction")]):
        _row(w, 2 + i, [rv, dd.isoformat(), ds])
    _widths(w, [6, 12, 30])
    wb.save(path)

def write_rfds_v3(s, path):
    wb = Workbook(); ws = wb.active; ws.title = "RFDS"
    ws["A1"] = f"{s['customer']} RFDS v3 | {s['site_id']} | SYNTHETIC SAMPLE"; ws["A1"].font = Font(bold=True, size=13)
    info = site_info(s)
    for i, (k, v) in enumerate(info):
        col = 1 + (i // 9) * 3; row = 3 + i % 9
        ws.cell(row=row, column=col, value=k.upper().replace(" ", "_")).font = Font(bold=True, size=8)
        ws.cell(row=row, column=col + 1, value=v).font = Font(size=8)
    cols = ["SECTOR", "AZ_DEG", "POS", "EQUIP_TYPE", "VENDOR", "MODEL", "ACTION", "BANDS", "TECH", "RC_FT", "MT_DEG", "ET_DEG", "PORTS", "FEEDS_ANT_POS"]
    _hdr(ws, 14, cols); r = 15
    for sc in s["sectors"]:
        for d in [d for d in s["devices"] if d["in_rfds"] and d["sector"] == sc["name"]]:
            c = M.CATALOG[d["key"]]; mt, el, eh = tilts(s, d)
            typ = {"air": "AAU", "radio": "RRU", "passive": "ANT"}[d["kind"]]
            act = {"Existing": "KEEP", "New": "ADD", "Remove": "REMOVE"}[d["status"]]
            _row(ws, r, [sc["name"], sc["az"], d["pos"], typ, c["mfr"], c["model"], act, c["bands"], c["tech"], d["rc"],
                         mt if typ != "RRU" else None, el if typ == "ANT" else (0 if typ == "AAU" else None),
                         c["ports"] or None, d["pos"] if typ == "RRU" else None]); r += 1
    _widths(ws, [10, 8, 6, 11, 11, 20, 9, 9, 8, 7, 7, 7, 7, 13])
    w = wb.create_sheet("BB_PWR"); _hdr(w, 1, ["EQUIP", "MODEL", "ACTION", "QTY"])
    for i, bb in enumerate(s["basebands"]): _row(w, 2 + i, ["BASEBAND", M.CATALOG[bb]["model"], "ADD", 1])
    w = wb.create_sheet("REV"); D = dates(s); _hdr(w, 1, ["REV", "DATE", "NOTE"])
    for i, (rv, dd, ds) in enumerate([("R1", D["rfds_r1"], "PRELIM"), ("R2", D["rfds_r2"], "UPDATE"), ("R3", D["rfds_r3"], "FINAL")]):
        _row(w, 2 + i, [rv, dd.isoformat(), ds])
    wb.save(path)

ST = getSampleStyleSheet()
SMALL = ParagraphStyle("s", parent=ST["Normal"], fontSize=7.5, leading=9)
def _tbl(data, widths=None, fs=7.5):
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", fs), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", fs),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    return t

def write_rfds_pdf(s, path):
    doc = SimpleDocTemplate(path, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30,
                            title=f"{s['site_id']} RFDS {s['rfds_rev']}")
    el = [Paragraph(f"<b>RF DATA SHEET - {s['customer']}</b>  |  {s['site_id']} {s['name']}  |  {s['rfds_rev']}", ST["Title"]),
          Paragraph("<font color='#C00000'>SYNTHETIC SAMPLE - NOT FOR CONSTRUCTION</font>", ST["Normal"]), Spacer(1, 6)]
    info = site_info(s)
    half = (len(info) + 1) // 2
    rows = [[a[0], str(a[1]), b[0] if b else "", str(b[1]) if b else ""] for a, b in zip(info[:half], info[half:] + [None])]
    el += [_tbl([["Field", "Value", "Field", "Value"]] + rows, [110, 220, 110, 220]), Spacer(1, 10)]
    for state, title in (("existing", "EXISTING CONFIGURATION"), ("final", "FINAL CONFIGURATION")):
        el.append(Paragraph(f"<b>{title}</b>", ST["Heading3"]))
        data = [["Sector", "Az", "Pos", "Action", "Antenna / AAU", "RC (ft)", "M-Tilt", "E-Tilt", "Radios (action)", "Bands"]]
        for sc, p, a, rads in rfds_positions(s, state):
            c = M.CATALOG[a["key"]]; mt, e1, e2 = tilts(s, a)
            rad = "Integrated" if a["kind"] == "air" else ", ".join(f"{M.CATALOG[x['key']]['model']} ({action(x)})" for x in rads)
            bands = c["bands"] if a["kind"] == "air" else "/".join(M.CATALOG[x["key"]]["bands"] for x in rads)
            data.append([SN[sc["name"]], sc["az"], p, action(a) if state == "final" else "Existing", c["model"], a["rc"], mt, f"{e1}/{e2}", rad, bands])
        el += [_tbl(data, [45, 30, 25, 45, 110, 40, 35, 40, 230, 80]), Spacer(1, 8)]
    load = sum(M.CATALOG[d["key"]]["power_w"] for d in s["devices"] if d["in_rfds"] and d["status"] == "New")
    el.append(Paragraph("<b>BASEBAND & POWER</b>", ST["Heading3"]))
    el.append(_tbl([["Item", "Model", "Action", "Qty"]] + [["Baseband", M.CATALOG[b]["model"], "Add", 1] for b in s["basebands"]]
                   + [["Router", "Router 6675", "Retain", 1], ["DC load added (W)", "", "", load]], [120, 150, 60, 60]))
    D = dates(s)
    el += [Spacer(1, 8), Paragraph("<b>REVISION HISTORY</b>", ST["Heading3"]),
           _tbl([["Rev", "Date", "Description"], ["R1", D["rfds_r1"].isoformat(), "Preliminary"],
                 ["R2", D["rfds_r2"].isoformat(), "Band plan update"], ["R3", D["rfds_r3"].isoformat(), "Final for construction"]], [40, 80, 200])]
    doc.build(el)

def write_ma_sa(s, path):
    D = dates(s)
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40, title=f"{s['site_id']} MA/SA")
    ma, sa = s["ma"], s["sa"]
    el = [Paragraph(f"<b>Mount Analysis Summary</b> - {s['site_id']} {s['name']}", ST["Title"]),
          Paragraph("Prepared by Meridian Structural Engineering (synthetic) for Ericsson. SYNTHETIC SAMPLE - NOT FOR CONSTRUCTION.", SMALL), Spacer(1, 8)]
    mount = "Roof sled mounts (non-penetrating)" if s["rooftop"] else ("New 12 ft sector frames (proposed)" if any(sc["frame"] == "new" for sc in s["sectors"]) else "Existing 12 ft sector frames")
    rows = [["Item", "Value"], ["Report", f"MA-{s['site_id']}-R0"], ["Date", D["ma"].isoformat()], ["Mount type", mount],
            ["Governing standard", "TIA-222-H"], ["Result", ma["result"]], ["Mount capacity (proposed loading)", f"{ma['capacity']}%"]]
    if "capacity_after" in ma: rows.append(["Capacity after modifications", f"{ma['capacity_after']}%"])
    el += [_tbl(rows, [200, 260], 9), Spacer(1, 10)]
    if s["ma_mods"]:
        el.append(Paragraph("<b>Required modifications</b>", ST["Heading3"]))
        el.append(_tbl([["Sector", "Modification", "Qty"]] + [[SN[a], f"Install mount reinforcement kit ({M.CATALOG[k]['model']}) per detail S-2", q] for a, k, q in s["ma_mods"]], [70, 330, 40], 9))
    else:
        el.append(Paragraph("No mount modifications required for the proposed loading.", ST["Normal"]))
    el += [Spacer(1, 10), Paragraph("<b>Proposed appurtenance loading analysed</b>", ST["Heading3"])]
    data = [["Sector", "Pos", "Equipment", "Action", "RC (ft)"]]
    for d in s["devices"]:
        if d["in_rfds"] and d["status"] in ("Existing", "New"):
            data.append([SN[d["sector"]], d["pos"], M.CATALOG[d["key"]]["model"], action(d), d["rc"]])
    el += [_tbl(data, [60, 30, 170, 60, 50], 8), PageBreak()]
    el += [Paragraph(f"<b>Structural Analysis Summary</b> - {s['site_id']} {s['name']}", ST["Title"]),
           Paragraph("Prepared by Meridian Structural Engineering (synthetic). SYNTHETIC SAMPLE - NOT FOR CONSTRUCTION.", SMALL), Spacer(1, 8),
           _tbl([["Item", "Value"], ["Report", f"SA-{s['site_id']}-R0"], ["Date", D["ma"].isoformat()], ["Structure", f"{s['height']} ft {s['structure']}"],
                 ["Owner", s["owner"]], ["Governing standard", "TIA-222-H, ASCE 7-16"], ["Result", sa["result"]],
                 ["Controlling member capacity", f"{sa['capacity']}%"]], [200, 260], 9), Spacer(1, 10)]
    if sa["result"] == "FAIL":
        el.append(Paragraph("<b>The structure does not have adequate capacity for the proposed loading. A structural modification design is required before installation.</b>", ST["Normal"]))
    else:
        el.append(Paragraph("The structure has adequate capacity for the proposed loading.", ST["Normal"]))
    doc.build(el)

def write_catalog(path):
    wb = Workbook(); ws = wb.active; ws.title = "Material Catalog"
    cols = ["Catalog Key", "Category", "Subcategory", "Manufacturer", "Model", "Manufacturer P/N", "Customer P/N", "Description",
            "UOM", "Supply (CFM/VFM)", "Customer Approved", "Unit Cost (USD)", "RF Ports", "Power (W)", "Bands", "Technology",
            "Height (m)", "Width (m)", "Depth (m)"]
    _hdr(ws, 1, cols)
    for i, c in enumerate(M.CATALOG.values()):
        _row(ws, 2 + i, [c["key"], c["category"], c["subcategory"], c["mfr"], c["model"], c["mfr_pn"], c["cust_pn"], c["desc"], c["uom"],
                         c["supply"], c["approved"], c["unit_cost"], c["ports"] or None, c["power_w"] or None, c["bands"], c["tech"],
                         c["h_m"] or None, c["w_m"] or None, c["d_m"] or None])
    _widths(ws, [11, 17, 20, 15, 22, 20, 11, 52, 6, 10, 10, 10, 7, 8, 9, 9, 8, 8, 8]); ws.freeze_panes = "B2"
    n = ws.max_row + 2
    ws.cell(row=n, column=1, value="SYNTHETIC: model names are representative; P/Ns, customer P/Ns, specs and costs are invented. Not for procurement.").font = Font(italic=True, color="C00000")
    wb.save(path)

def write_rules(path):
    wb = Workbook(); ws = wb.active; ws.title = "BOM Rules"
    _hdr(ws, 1, ["Rule ID", "Trigger", "Child item (catalog key)", "Quantity", "Notes"])
    for i, r in enumerate(M.RULES): _row(ws, 2 + i, list(r))
    _widths(ws, [11, 46, 22, 36, 50]); ws.freeze_panes = "A2"
    w = wb.create_sheet("Trunk Lengths"); _hdr(w, 1, ["Catalog step (ft)", "Catalog key"])
    for i, L in enumerate(M.TRUNK_STEPS): _row(w, 2 + i, [L, f"HYB{L}"])
    w.cell(row=10, column=1, value="Required length = 1.10 x (vertical + horizontal run). Tower: vertical = highest new radio/AAU rad center + 5 ft. Rooftop: vertical = riser run.").font = Font(italic=True)
    wb.save(path)

def write_bom_xlsx(path, sheets, title_note):
    wb = Workbook(); first = True
    for name, rows in sheets:
        ws = wb.active if first else wb.create_sheet(); first = False
        ws.title = name[:31]
        _hdr(ws, 1, M.COLS)
        for i, r in enumerate(rows): _row(ws, 2 + i, r)
        _widths(ws, [10, 7, 8, 9, 9, 17, 14, 20, 18, 10, 44, 6, 7, 7, 7, 11, 8, 26, 10, 8, 13]); ws.freeze_panes = "C2"
        ws.cell(row=len(rows) + 3, column=1, value=title_note).font = Font(italic=True, color="C00000")
    wb.save(path)

def sf_id(prefix, key):
    alpha = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"
    n = int(hashlib.sha1(key.encode()).hexdigest(), 16)
    body = "".join(alpha[(n >> (6 * i)) % 62] for i in range(15 - len(prefix)))
    return (prefix + body + "AAA")[:18]

def write_sitetracker(S, folder, doc_index):
    def w(name, cols, rows):
        with open(f"{folder}/{name}", "w", newline="") as f:
            cw = csv.writer(f); cw.writerow(cols); cw.writerows(rows)
    w("Site.csv", ["Id", "Name", "Site_Name__c", "Market__c", "Customer__c", "Street__c", "City__c", "State__c", "Zip__c", "Latitude__c",
                   "Longitude__c", "Structure_Type__c", "Structure_Height_ft__c", "Structure_Owner__c", "Site_Status__c"],
      [[sf_id("a0S", s["site_id"]), s["site_id"], s["name"], s["market"], s["customer"], s["addr"], s["city"], s["state"], s["zip"],
        s["lat"], s["lon"], s["structure"], s["height"], s["owner"], "On Air"] for s in S])
    lines = {"T1": "Site Build BOM Scoping", "T2": "Network Deployment Scoping", "T3": "Site Build BOM Scoping"}
    prj, ms = [], []
    for s in S:
        D = dates(s); pid = sf_id("a0P", s["site_id"] + "P")
        prj.append([pid, f"{s['site_id']}-{s['scope']}-2026", sf_id("a0S", s["site_id"]), s["project"], lines[s["scope"]],
                    D["assigned"].isoformat(), "REV 0 - Preliminary", "REV 0", s["cx_sp"] + " (synthetic)",
                    (D["assigned"] + dt.timedelta(days=21)).isoformat(), "", ""])
        steps = [("RFDS Final", D["rfds_r3"], True), ("Drone Survey Complete", D["flight"], True), ("CDs Final (REV 1)", D["cd"], True),
                 ("MA/SA Complete", D["ma"], True), ("BOM REV 0 Created", D["rev0"], True), ("Scoping Assigned", D["assigned"], True),
                 ("BOM REV 1 Submitted", D["assigned"] + dt.timedelta(days=7), False), ("CX SP Handshake", D["assigned"] + dt.timedelta(days=12), False),
                 ("Final BOM Approval (FBA)", D["assigned"] + dt.timedelta(days=18), False), ("Scoping Complete", D["assigned"] + dt.timedelta(days=21), False)]
        for i, (n, d, done) in enumerate(steps):
            ms.append([sf_id("a0M", s["site_id"] + n), pid, n, d.isoformat(), d.isoformat() if done else "", "Complete" if done else "Open", i + 1])
    w("Project.csv", ["Id", "Name", "Site__c", "Project_Type__c", "Scoping_Service_Line__c", "Scoping_Assigned_Date__c", "BOM_Status__c",
                      "Current_BOM_Rev__c", "CX_Service_Provider__c", "Scoping_Complete_Forecast__c", "Scoping_Complete_Actual__c", "FBA_Date__c"], prj)
    w("Milestone.csv", ["Id", "Project__c", "Name", "Forecast_Date__c", "Actual_Date__c", "Status__c", "Sequence__c"], ms)
    w("Document.csv", ["Id", "Project__c", "Document_Type__c", "Title", "File_Path__c", "File_Format__c", "Revision__c", "Uploaded_Date__c"],
      [[sf_id("069", r[1]), sf_id("a0P", r[0] + "P"), r[2], r[1].split("/")[-1], r[1], r[3], r[4], r[5]] for r in doc_index])
