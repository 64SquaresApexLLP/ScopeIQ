"""Document content comparison and diffing service.

Detects:
1. NEW data: file does not match any existing site document.
2. DUPLICATE data: file content is identical to an existing document, even if filenames differ or hashes differ (e.g. re-saved spreadsheets).
   Prompts: "Duplicate data detected. Do you want to replace it with the existing file?"
3. UPDATED data: file contains old data from an existing document with new additions (incremental update).
   Prompts: "Replace old data with this updated data?" and shows "what is new in the new file".
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from scopeiq.common.logging import get_logger
from scopeiq.config import get_settings

log = get_logger(__name__)


@dataclass
class DiffResult:
    status: str  # "NEW" | "DUPLICATE" | "UPDATED"
    match_type: str  # "EXACT_HASH" | "CONTENT_IDENTICAL" | "CONTENT_UPDATED" | "NONE"
    existing_file: str | None = None
    existing_path: str | None = None
    existing_source: str | None = None  # "DOCUMENT" | "UPLOAD" | "SITE_FOLDER" | "SHAREPOINT"
    doc_type: str | None = None
    file_hash: str | None = None
    existing_hash: str | None = None
    message: str | None = None
    warning: str | None = None
    added_count: int = 0
    summary: str | None = None
    additions: list[str] = field(default_factory=list)
    details: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "is_duplicate": self.status == "DUPLICATE",
            "is_updated": self.status == "UPDATED",
            "is_new": self.status == "NEW",
            "match_type": self.match_type,
            "existing_file": self.existing_file,
            "duplicate_of": self.existing_file if self.status == "DUPLICATE" else None,
            "updated_from": self.existing_file if self.status == "UPDATED" else None,
            "existing_path": self.existing_path,
            "existing_source": self.existing_source,
            "doc_type": self.doc_type,
            "file_hash": self.file_hash,
            "existing_hash": self.existing_hash,
            "message": self.message,
            "warning": self.warning,
            "diff": {
                "summary": self.summary,
                "added_count": self.added_count,
                "additions": self.additions[:50],  # cap at 50 for clean JSON response
                "total_additions": self.added_count,
            } if self.status == "UPDATED" else None,
            "details": self.details,
        }


def _normalize_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def extract_excel_sheets(data: bytes) -> dict[str, list[tuple[str, ...]]]:
    """Extract non-empty rows per sheet from an Excel file in memory."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        sheets: dict[str, list[tuple[str, ...]]] = {}
        for name in wb.sheetnames:
            ws = wb[name]
            rows: list[tuple[str, ...]] = []
            for r in ws.iter_rows(values_only=True):
                norm = tuple(_normalize_cell(c) for c in r)
                # Trim trailing empty values
                while norm and not norm[-1]:
                    norm = norm[:-1]
                if any(norm):
                    rows.append(norm)
            if rows:
                sheets[name] = rows
        return sheets
    except Exception as ex:
        log.warning("failed to parse Excel content: %s", ex)
        return {}


def extract_csv_rows(data: bytes) -> list[tuple[str, ...]]:
    """Extract non-empty rows from CSV data."""
    try:
        text = data.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = []
        for r in reader:
            norm = tuple(c.strip() for c in r)
            while norm and not norm[-1]:
                norm = norm[:-1]
            if any(norm):
                rows.append(norm)
        return rows
    except Exception as ex:
        log.warning("failed to parse CSV content: %s", ex)
        return []


def extract_json_data(data: bytes) -> Any:
    """Parse JSON data."""
    try:
        return json.loads(data.decode("utf-8", errors="replace"))
    except Exception:
        return None


def extract_pdf_text(data: bytes) -> list[str]:
    """Extract text per page from PDF bytes."""
    pages: list[str] = []
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        for p in reader.pages:
            t = (p.extract_text() or "").strip()
            if t:
                pages.append(t)
    except Exception:
        try:
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            for page in doc:
                t = page.get_text().strip()
                if t:
                    pages.append(t)
        except Exception as ex:
            log.warning("failed to parse PDF text: %s", ex)
    return pages


def compare_excel_data(new_sheets: dict[str, list[tuple[str, ...]]],
                       old_sheets: dict[str, list[tuple[str, ...]]]) -> tuple[str, list[str], int]:
    """
    Compares two parsed Excel workbooks.
    Returns (status, additions, added_count):
    status: "DUPLICATE" | "UPDATED" | "DIFFERENT"
    """
    if not new_sheets or not old_sheets:
        return "DIFFERENT", [], 0

    # 1. Exact match check across all sheets and rows
    if new_sheets == old_sheets:
        return "DUPLICATE", [], 0

    # 2. Check for subset / incremental update
    all_old_rows: set[tuple[str, ...]] = set()
    for s_rows in old_sheets.values():
        all_old_rows.update(s_rows)

    all_new_rows: set[tuple[str, ...]] = set()
    for s_rows in new_sheets.values():
        all_new_rows.update(s_rows)

    if not all_old_rows:
        return "DIFFERENT", [], 0

    # How many of the old rows exist in the new workbook?
    common_rows = all_old_rows.intersection(all_new_rows)
    old_coverage = len(common_rows) / len(all_old_rows)

    # If at least 80% of old rows are retained, and new rows were added:
    if old_coverage >= 0.80 and len(all_new_rows) > len(all_old_rows):
        additions: list[str] = []
        for sname, nrows in new_sheets.items():
            orows = set(old_sheets.get(sname, []))
            for i, r in enumerate(nrows, 1):
                if r not in orows and any(r):
                    clean = " | ".join(c for c in r if c)
                    additions.append(f"[{sname}] Row {i}: {clean}")
        if additions:
            return "UPDATED", additions, len(additions)

    # Check sheet by sheet if old sheets are subset
    shared_sheets = set(old_sheets.keys()).intersection(set(new_sheets.keys()))
    if shared_sheets:
        additions = []
        is_superset = True
        for sname in old_sheets.keys():
            if sname not in new_sheets:
                is_superset = False
                break
            orows = set(old_sheets[sname])
            nrows = set(new_sheets[sname])
            if len(orows.intersection(nrows)) / max(len(orows), 1) < 0.80:
                is_superset = False
                break
            for i, r in enumerate(new_sheets[sname], 1):
                if r not in orows and any(r):
                    clean = " | ".join(c for c in r if c)
                    additions.append(f"[{sname}] Row {i}: {clean}")
        if is_superset and additions:
            return "UPDATED", additions, len(additions)

    return "DIFFERENT", [], 0


def compare_csv_data(new_rows: list[tuple[str, ...]],
                     old_rows: list[tuple[str, ...]]) -> tuple[str, list[str], int]:
    """Compares two parsed CSV file row sets."""
    if not new_rows or not old_rows:
        return "DIFFERENT", [], 0
    if new_rows == old_rows:
        return "DUPLICATE", [], 0

    old_set = set(old_rows)
    new_set = set(new_rows)
    common = old_set.intersection(new_set)
    coverage = len(common) / max(len(old_set), 1)

    if coverage >= 0.80 and len(new_rows) > len(old_rows):
        additions = [f"Row {i}: {' | '.join(c for c in r if c)}" for i, r in enumerate(new_rows, 1) if r not in old_set]
        if additions:
            return "UPDATED", additions, len(additions)

    return "DIFFERENT", [], 0


def compare_json_data(new_obj: Any, old_obj: Any) -> tuple[str, list[str], int]:
    """Compares two JSON objects."""
    if new_obj == old_obj:
        return "DUPLICATE", [], 0

    additions = []
    if isinstance(new_obj, dict) and isinstance(old_obj, dict):
        new_keys = set(new_obj.keys()) - set(old_obj.keys())
        for k in sorted(new_keys):
            additions.append(f"New property '{k}': {str(new_obj[k])[:80]}")
        # Check shared keys that are lists or dicts
        for k in set(new_obj.keys()).intersection(set(old_obj.keys())):
            nv, ov = new_obj[k], old_obj[k]
            if isinstance(nv, list) and isinstance(ov, list) and len(nv) > len(ov):
                for item in nv:
                    if item not in ov:
                        additions.append(f"In '{k}': added {str(item)[:80]}")
        if additions and len(set(old_obj.keys()).intersection(set(new_obj.keys()))) >= len(old_obj) * 0.7:
            return "UPDATED", additions, len(additions)

    if isinstance(new_obj, list) and isinstance(old_obj, list):
        if len(new_obj) > len(old_obj) and all(item in new_obj for item in old_obj):
            for item in new_obj:
                if item not in old_obj:
                    additions.append(f"Added item: {str(item)[:80]}")
            return "UPDATED", additions, len(additions)

    return "DIFFERENT", [], 0


def compare_text_data(new_text: str, old_text: str) -> tuple[str, list[str], int]:
    """Compares two plain text files line-by-line."""
    new_lines = [l.strip() for l in new_text.splitlines() if l.strip()]
    old_lines = [l.strip() for l in old_text.splitlines() if l.strip()]
    if not new_lines or not old_lines:
        return "DIFFERENT", [], 0
    if new_lines == old_lines:
        return "DUPLICATE", [], 0

    old_set = set(old_lines)
    common = old_set.intersection(set(new_lines))
    if len(common) / max(len(old_set), 1) >= 0.80 and len(new_lines) > len(old_lines):
        additions = [f"Line {i}: {l}" for i, l in enumerate(new_lines, 1) if l not in old_set]
        if additions:
            return "UPDATED", additions, len(additions)

    return "DIFFERENT", [], 0


def compare_pdf_data(new_pages: list[str], old_pages: list[str]) -> tuple[str, list[str], int]:
    """Compares two PDFs by extracted text pages."""
    if not new_pages or not old_pages:
        return "DIFFERENT", [], 0
    if new_pages == old_pages:
        return "DUPLICATE", [], 0

    new_full = " ".join(new_pages)
    old_full = " ".join(old_pages)
    if new_full == old_full:
        return "DUPLICATE", [], 0

    if len(new_pages) > len(old_pages) and all(p in new_pages for p in old_pages):
        additions = [f"Page {i + 1} added: {p[:120]}..." for i, p in enumerate(new_pages) if p not in old_pages]
        return "UPDATED", additions, len(additions)

    if old_full in new_full and len(new_full) > len(old_full) + 15:
        added_text = new_full.replace(old_full, "").strip()
        return "UPDATED", [f"Added text: {added_text[:150]}"], 1

    return "DIFFERENT", [], 0


def get_existing_site_files(site_id: str, repo: Any | None = None) -> list[dict[str, Any]]:
    """
    Gathers all known documents and files for this site from:
    1. CORE.DOCUMENT table
    2. CORE.UPLOAD table
    3. Disk: paths.upload_dir / site_id
    4. Disk: paths.sites_dir / site_id
    5. SharePoint library (if local)
    """
    found: dict[str, dict[str, Any]] = {}  # key by file name or rel_path
    st = get_settings()

    # 1. DB tables
    if repo:
        try:
            for d in repo.select("CORE.DOCUMENT", {"SITE_ID": site_id}):
                name = d.get("FILE_NAME")
                if not name:
                    continue
                rel = d.get("REL_PATH") or ""
                prefix, _, rest = rel.partition("/")
                base = st.path("paths.sites_dir") if prefix == "sites" else st.path("paths.upload_dir")
                p = base / rest
                found[name] = {
                    "name": name,
                    "path": p if p.is_file() else None,
                    "doc_type": d.get("DOC_TYPE"),
                    "hash": d.get("FILE_HASH"),
                    "source": "DOCUMENT",
                }
        except Exception:
            pass

        try:
            for u in repo.select("CORE.UPLOAD", {"SITE_ID": site_id}):
                name = u.get("FILE_NAME")
                if not name or u.get("STATUS") == "SUPERSEDED":
                    continue
                p = Path(st.path("paths.upload_dir")) / site_id / (u.get("STAGE_PATH") or "").split("/")[-2] / name if "/" in (u.get("STAGE_PATH") or "") else None
                found[name] = {
                    "name": name,
                    "path": p if p and p.is_file() else None,
                    "doc_type": u.get("DOC_TYPE"),
                    "hash": u.get("FILE_HASH"),
                    "source": "UPLOAD",
                }
        except Exception:
            pass

    # 2. Disk scan: upload_dir / site_id
    try:
        up_root = Path(st.path("paths.upload_dir")) / site_id
        if up_root.is_dir():
            for p in up_root.rglob("*"):
                if p.is_file() and not p.name.startswith("."):
                    if p.name not in found or not found[p.name]["path"]:
                        found[p.name] = {
                            "name": p.name,
                            "path": p,
                            "doc_type": None,
                            "hash": None,
                            "source": "UPLOAD",
                        }
    except Exception:
        pass

    # 3. Disk scan: sites_dir / site_id
    try:
        site_root = Path(st.path("paths.sites_dir")) / site_id
        if site_root.is_dir():
            for p in site_root.rglob("*"):
                if p.is_file() and not p.name.startswith("."):
                    if p.name not in found or not found[p.name]["path"]:
                        found[p.name] = {
                            "name": p.name,
                            "path": p,
                            "doc_type": None,
                            "hash": None,
                            "source": "SITE_FOLDER",
                        }
    except Exception:
        pass

    # 4. SharePoint local library
    try:
        sp_local = st.path("sharepoint.local_root")
        sp_lib = st.get("sharepoint.library_root", "Ericsson-Scoping/Sites")
        sp_site = Path(sp_local) / sp_lib / site_id
        if sp_site.is_dir():
            for p in sp_site.rglob("*"):
                if p.is_file() and not p.name.startswith("."):
                    if p.name not in found or not found[p.name]["path"]:
                        found[p.name] = {
                            "name": p.name,
                            "path": p,
                            "doc_type": None,
                            "hash": None,
                            "source": "SHAREPOINT",
                        }
    except Exception:
        pass

    return list(found.values())


def inspect_file_against_site(
    site_id: str,
    new_filename: str,
    new_data: bytes,
    doc_type: str | None = None,
    repo: Any | None = None,
    exclude_filename: str | None = None,
) -> DiffResult:
    """
    Main inspection engine:
    Takes an uploaded or picked file, and checks it against all existing files for the site.
    Returns DiffResult with status: "NEW", "DUPLICATE", or "UPDATED".
    """
    if not new_data:
        return DiffResult(status="NEW", match_type="NONE", message="File is empty")

    new_hash = hashlib.sha256(new_data).hexdigest()
    ext = Path(new_filename).suffix.lower()

    # Pre-parse new data based on extension
    new_excel = extract_excel_sheets(new_data) if ext in (".xlsx", ".xls", ".xlsm") else None
    new_csv = extract_csv_rows(new_data) if ext == ".csv" else None
    new_json = extract_json_data(new_data) if ext == ".json" else None
    new_pdf = extract_pdf_text(new_data) if ext == ".pdf" else None
    new_text = new_data.decode("utf-8", errors="ignore") if ext in (".txt", ".log", ".md") else None

    existing_files = get_existing_site_files(site_id, repo=repo)

    # Prioritize files of the same doc_type or matching extension
    matching_candidates: list[dict[str, Any]] = []
    other_candidates: list[dict[str, Any]] = []

    for f in existing_files:
        if exclude_filename and f["name"] == exclude_filename:
            continue
        f_ext = Path(f["name"]).suffix.lower()
        if f.get("doc_type") == doc_type or f_ext == ext:
            matching_candidates.append(f)
        else:
            other_candidates.append(f)

    all_candidates = matching_candidates + other_candidates

    best_update_match: DiffResult | None = None

    for cand in all_candidates:
        cand_name = cand["name"]
        if exclude_filename and cand_name == exclude_filename:
            continue
        cand_path = cand["path"]
        cand_hash = cand.get("hash")
        cand_source = cand.get("source") or "SITE_FOLDER"

        # 1. Exact SHA-256 hash match
        if cand_hash and cand_hash.lower() == new_hash.lower():
            return DiffResult(
                status="DUPLICATE",
                match_type="EXACT_HASH",
                existing_file=cand_name,
                existing_path=str(cand_path) if cand_path else None,
                existing_source=cand_source,
                doc_type=cand.get("doc_type") or doc_type,
                file_hash=new_hash,
                existing_hash=cand_hash,
                message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}'. "
                        f"Do you want to replace it with the existing file?",
                warning=f"Duplicate document detected! Identical SHA-256 hash match to '{cand_name}'."
            )

        # If we have the file on disk, read it for content comparison
        if not cand_path or not Path(cand_path).is_file():
            continue

        try:
            cand_bytes = Path(cand_path).read_bytes()
        except Exception:
            continue

        cand_actual_hash = hashlib.sha256(cand_bytes).hexdigest()
        if cand_actual_hash == new_hash:
            return DiffResult(
                status="DUPLICATE",
                match_type="EXACT_HASH",
                existing_file=cand_name,
                existing_path=str(cand_path),
                existing_source=cand_source,
                doc_type=cand.get("doc_type") or doc_type,
                file_hash=new_hash,
                existing_hash=cand_actual_hash,
                message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}'. "
                        f"Do you want to replace it with the existing file?",
                warning=f"Duplicate document detected! Identical SHA-256 hash match to '{cand_name}'."
            )

        # 2. Content comparison by file type
        c_ext = Path(cand_name).suffix.lower()

        # Excel
        if new_excel is not None and c_ext in (".xlsx", ".xls", ".xlsm"):
            old_excel = extract_excel_sheets(cand_bytes)
            status, additions, added_count = compare_excel_data(new_excel, old_excel)
            if status == "DUPLICATE":
                return DiffResult(
                    status="DUPLICATE",
                    match_type="CONTENT_IDENTICAL",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}' "
                            f"(even though file names or binary formats differ). Do you want to replace it with the existing file?",
                    warning=f"Duplicate data detected! All spreadsheet rows match existing file '{cand_name}'."
                )
            elif status == "UPDATED" and not best_update_match:
                summary = f"{added_count} new row{'s' if added_count > 1 else ''} added across sheets"
                best_update_match = DiffResult(
                    status="UPDATED",
                    match_type="CONTENT_UPDATED",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Updated data detected! This file contains existing data from '{cand_name}' with {added_count} new addition{'s' if added_count > 1 else ''}. "
                            f"Do you want to replace old data with this updated data?",
                    warning=f"Updated data: Contains previous content from '{cand_name}' with new additions.",
                    added_count=added_count,
                    summary=summary,
                    additions=additions,
                )

        # CSV
        elif new_csv is not None and c_ext == ".csv":
            old_csv = extract_csv_rows(cand_bytes)
            status, additions, added_count = compare_csv_data(new_csv, old_csv)
            if status == "DUPLICATE":
                return DiffResult(
                    status="DUPLICATE",
                    match_type="CONTENT_IDENTICAL",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}'. "
                            f"Do you want to replace it with the existing file?",
                    warning=f"Duplicate data detected! All CSV rows match existing file '{cand_name}'."
                )
            elif status == "UPDATED" and not best_update_match:
                best_update_match = DiffResult(
                    status="UPDATED",
                    match_type="CONTENT_UPDATED",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Updated data detected! This file contains existing data from '{cand_name}' with {added_count} new row{'s' if added_count > 1 else ''}. "
                            f"Do you want to replace old data with this updated data?",
                    warning=f"Updated data: Contains previous CSV rows with {added_count} new row additions.",
                    added_count=added_count,
                    summary=f"{added_count} new row{'s' if added_count > 1 else ''} added",
                    additions=additions,
                )

        # JSON
        elif new_json is not None and c_ext == ".json":
            old_json = extract_json_data(cand_bytes)
            if old_json is not None:
                status, additions, added_count = compare_json_data(new_json, old_json)
                if status == "DUPLICATE":
                    return DiffResult(
                        status="DUPLICATE",
                        match_type="CONTENT_IDENTICAL",
                        existing_file=cand_name,
                        existing_path=str(cand_path),
                        existing_source=cand_source,
                        doc_type=cand.get("doc_type") or doc_type,
                        file_hash=new_hash,
                        existing_hash=cand_actual_hash,
                        message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}'. "
                                f"Do you want to replace it with the existing file?",
                        warning=f"Duplicate JSON data matches existing file '{cand_name}'."
                    )
                elif status == "UPDATED" and not best_update_match:
                    best_update_match = DiffResult(
                        status="UPDATED",
                        match_type="CONTENT_UPDATED",
                        existing_file=cand_name,
                        existing_path=str(cand_path),
                        existing_source=cand_source,
                        doc_type=cand.get("doc_type") or doc_type,
                        file_hash=new_hash,
                        existing_hash=cand_actual_hash,
                        message=f"Updated data detected! This file contains existing data from '{cand_name}' with {added_count} new addition{'s' if added_count > 1 else ''}. "
                                f"Do you want to replace old data with this updated data?",
                        warning=f"Updated JSON: Contains previous data with {added_count} new additions.",
                        added_count=added_count,
                        summary=f"{added_count} new item{'s' if added_count > 1 else ''} added",
                        additions=additions,
                    )

        # Text
        elif new_text is not None and c_ext in (".txt", ".log", ".md"):
            old_text = cand_bytes.decode("utf-8", errors="ignore")
            status, additions, added_count = compare_text_data(new_text, old_text)
            if status == "DUPLICATE":
                return DiffResult(
                    status="DUPLICATE",
                    match_type="CONTENT_IDENTICAL",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}'. "
                            f"Do you want to replace it with the existing file?",
                    warning=f"Duplicate text content matches existing file '{cand_name}'."
                )
            elif status == "UPDATED" and not best_update_match:
                best_update_match = DiffResult(
                    status="UPDATED",
                    match_type="CONTENT_UPDATED",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Updated data detected! This file contains existing data from '{cand_name}' with {added_count} new line{'s' if added_count > 1 else ''}. "
                            f"Do you want to replace old data with this updated data?",
                    warning=f"Updated text: Contains previous text with {added_count} new lines.",
                    added_count=added_count,
                    summary=f"{added_count} new line{'s' if added_count > 1 else ''} added",
                    additions=additions,
                )

        # PDF
        elif new_pdf is not None and c_ext == ".pdf":
            old_pdf = extract_pdf_text(cand_bytes)
            status, additions, added_count = compare_pdf_data(new_pdf, old_pdf)
            if status == "DUPLICATE":
                return DiffResult(
                    status="DUPLICATE",
                    match_type="CONTENT_IDENTICAL",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Duplicate data detected! The content of this file is identical to existing file '{cand_name}'. "
                            f"Do you want to replace it with the existing file?",
                    warning=f"Duplicate PDF text matches existing file '{cand_name}'."
                )
            elif status == "UPDATED" and not best_update_match:
                best_update_match = DiffResult(
                    status="UPDATED",
                    match_type="CONTENT_UPDATED",
                    existing_file=cand_name,
                    existing_path=str(cand_path),
                    existing_source=cand_source,
                    doc_type=cand.get("doc_type") or doc_type,
                    file_hash=new_hash,
                    existing_hash=cand_actual_hash,
                    message=f"Updated data detected! This file contains existing data from '{cand_name}' with {added_count} new section{'s' if added_count > 1 else ''}. "
                            f"Do you want to replace old data with this updated data?",
                    warning=f"Updated PDF: Contains previous document pages with new additions.",
                    added_count=added_count,
                    summary=f"{added_count} new section{'s' if added_count > 1 else ''}/page{'s' if added_count > 1 else ''} added",
                    additions=additions,
                )

    if best_update_match:
        return best_update_match

    # If no duplicate or update match found, file data is brand NEW!
    return DiffResult(
        status="NEW",
        match_type="NONE",
        file_hash=new_hash,
        doc_type=doc_type,
        message="New file data. Ready to store.",
    )
