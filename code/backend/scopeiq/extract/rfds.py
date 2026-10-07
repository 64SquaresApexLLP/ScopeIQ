"""RFDS extraction for the three layouts in circulation: XLSX template v2, XLSX template v3 and PDF export.

Output: site info, revision + history, EXISTING configuration and FINAL configuration as ConfigItems, the
baseband list and the added DC load. Every item carries a SourceRef (sheet/row or page) for traceability.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from scopeiq.common.errors import ExtractionError
from scopeiq.common.logging import get_logger, log_call
from scopeiq.domain import ConfigItem, ExtractedField, SourceRef, sector_code
from scopeiq.extract.pdftext import pdf_pages_text
from scopeiq.reference.catalog import Catalog

log = get_logger(__name__)

ACTION_MAP = {"RETAIN": "Retain", "KEEP": "Retain", "EXISTING": "Existing", "ADD": "New", "NEW": "New", "SWAP": "New",
              "REMOVE": "Remove", "REPLACE": "New", "RELOCATE": "Relocate"}


class RfdsResult:
    def __init__(self):
        self.site: dict = {}
        self.revision = ""
        self.date = ""
        self.history: list[dict] = []
        self.existing: list[ConfigItem] = []
        self.final: list[ConfigItem] = []
        self.basebands: list[dict] = []
        self.dc_load_w: float | None = None
        self.fields: list[ExtractedField] = []
        self.template = ""


def _f(v) -> float | None:
    try:
        return float(v) if v not in (None, "", "-") else None
    except (TypeError, ValueError):
        return None


class _Builder:
    def __init__(self, site_id: str, doc_id: str, catalog: Catalog, method: str, conf: float, rev: str):
        self.site_id, self.doc_id, self.catalog, self.method, self.conf, self.rev = site_id, doc_id, catalog, method, conf, rev

    def item(self, source: str, sector, pos, model: str, status: str, *, rc=None, az=None, tilt=None, bands="", page="",
             kind_hint: str | None = None, feeds=None) -> ConfigItem:
        kinds = (kind_hint,) if kind_hint else None
        cat, score = self.catalog.resolve(model, kinds=kinds)
        return ConfigItem(site_id=self.site_id, source=source, sector=sector_code(sector), position=int(pos) if pos not in (None, "") else None,
                          kind=cat.kind if cat else (kind_hint or "unknown"), catalog_key=cat.key if cat else None,
                          model_text=str(model).strip(), status=status, rad_center_ft=_f(rc), azimuth_deg=_f(az), mech_tilt=_f(tilt),
                          bands=str(bands or ""), ref=SourceRef(self.doc_id, "RFDS", self.rev, page),
                          confidence=round(self.conf * (score if cat else 0.3), 3), feeds_position=int(feeds) if feeds not in (None, "") else None)


def _radios(text: str) -> list[tuple[str, str]]:
    """'Radio 4449 B5 B12 (Retain), Radio 4460 B2 B66 (Add)' -> [(model, action)]"""
    out = []
    for part in [p.strip() for p in str(text or "").split(",") if p.strip()]:
        if part.lower().startswith("integrated"):
            continue
        m = re.match(r"(.+?)\s*\((\w+)\)\s*$", part)
        out.append((m.group(1).strip(), m.group(2)) if m else (part, ""))
    return out


@log_call()
def extract_rfds(path: Path, *, site_id: str, doc_id: str, catalog: Catalog) -> RfdsResult:
    path = Path(path)
    try:
        if path.suffix.lower() in (".xlsx", ".xlsm"):
            wb = openpyxl.load_workbook(path, data_only=True)
            if "Final Config" in wb.sheetnames:
                res = _xlsx_v2(wb, site_id, doc_id, catalog)
            elif "RFDS" in wb.sheetnames:
                res = _xlsx_v3(wb, site_id, doc_id, catalog)
            else:
                raise ExtractionError(f"Unknown RFDS workbook layout in {path.name}", details={"sheets": wb.sheetnames})
        elif path.suffix.lower() == ".pdf":
            res = _pdf(path, site_id, doc_id, catalog)
        else:
            raise ExtractionError(f"Unsupported RFDS format {path.suffix}")
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"RFDS {path.name} could not be read: {exc}", cause=exc, details={"file": path.name}) from exc
    for it in res.existing + res.final:
        res.fields.append(ExtractedField(doc_id, site_id, it.ref.page if it.ref else "", f"{it.source}:{it.sector}{it.position}:{it.kind}",
                                         it.model_text, it.confidence, res.template, needs_review=it.catalog_key is None))
    log.info("RFDS %s %s read: %d existing, %d final items", path.name, res.revision, len(res.existing), len(res.final),
             extra={"site_id": site_id, "template": res.template})
    return res


# ------------------------------------------------------------------ XLSX template v2
def _xlsx_v2(wb, site_id, doc_id, catalog) -> RfdsResult:
    res = RfdsResult()
    res.template = "XLSX_V2"
    info = {str(r[0]).strip(): r[1] for r in wb["Site Info"].iter_rows(values_only=True) if r and r[0]}
    res.site = info
    res.revision, res.date = str(info.get("RFDS Revision", "")), str(info.get("RFDS Date", ""))
    b = _Builder(site_id, doc_id, catalog, "XLSX", 0.99, res.revision)
    hdr_ex = None
    for i, r in enumerate(wb["Existing Config"].iter_rows(values_only=True), 1):
        if r and r[0] == "Sector":
            hdr_ex = {h: k for k, h in enumerate(r)}
            continue
        if not hdr_ex or not r or not r[0]:
            continue
        g = lambda c: r[hdr_ex[c]] if c in hdr_ex else None
        page = f"Existing Config!R{i}"
        res.existing.append(b.item("RFDS_EXISTING", g("Sector"), g("Ant Position"), g("Antenna Model"), "Existing", rc=g("Rad Center (ft)"),
                                   az=g("Azimuth (deg)"), tilt=g("Mech Tilt (deg)"), page=page))
        bands = str(g("Radio Bands") or "").split(",")
        for k, model in enumerate([m.strip() for m in str(g("Radios") or "").split(",") if m.strip()]):
            res.existing.append(b.item("RFDS_EXISTING", g("Sector"), g("Ant Position"), model, "Existing", rc=g("Rad Center (ft)"),
                                       az=g("Azimuth (deg)"), bands=bands[k].strip() if k < len(bands) else "", page=page,
                                       kind_hint="radio", feeds=g("Ant Position")))
    hdr = None
    for i, r in enumerate(wb["Final Config"].iter_rows(values_only=True), 1):
        if r and r[0] == "Sector":
            hdr = {h: k for k, h in enumerate(r)}
            continue
        if not hdr or not r or not r[0]:
            continue
        g = lambda c: r[hdr[c]] if c in hdr else None
        page = f"Final Config!R{i}"
        act = ACTION_MAP.get(str(g("Action")).upper(), str(g("Action")))
        res.final.append(b.item("RFDS_FINAL", g("Sector"), g("Ant Position"), g("Antenna / AAU Model"), act, rc=g("Rad Center (ft)"),
                                az=g("Azimuth (deg)"), tilt=g("Mech Tilt (deg)"), bands=g("Bands"), page=page))
        for n in (1, 2):
            model = g(f"Radio {n}")
            if not model or str(model).lower().startswith("integrated"):
                continue
            ract = ACTION_MAP.get(str(g(f"Radio {n} Action")).upper(), "Retain")
            res.final.append(b.item("RFDS_FINAL", g("Sector"), g("Ant Position"), model, ract, rc=g("Rad Center (ft)"),
                                    az=g("Azimuth (deg)"), page=page, kind_hint="radio", feeds=g("Ant Position")))
    for r in wb["Baseband & Power"].iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        if str(r[0]).startswith("DC load"):
            res.dc_load_w = _f(r[3])
        elif str(r[0]).lower() == "baseband":
            res.basebands.append({"model": r[1], "action": ACTION_MAP.get(str(r[2]).upper(), r[2]), "qty": int(r[3] or 1)})
    if "Revision History" in wb.sheetnames:
        res.history = [{"rev": r[0], "date": str(r[1]), "note": r[2]} for r in wb["Revision History"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    return res


# ------------------------------------------------------------------ XLSX template v3
def _xlsx_v3(wb, site_id, doc_id, catalog) -> RfdsResult:
    res = RfdsResult()
    res.template = "XLSX_V3"
    ws = wb["RFDS"]
    info, hdr = {}, None
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows, 1):
        if r and r[0] == "SECTOR":
            hdr = {h: k for k, h in enumerate(r) if h}
            continue
        if hdr is None:
            for k in (0, 3):
                if r and len(r) > k + 1 and r[k]:
                    info[str(r[k]).strip()] = r[k + 1]
            continue
        if not r or not r[0]:
            continue
    res.site = {k.replace("_", " ").title(): v for k, v in info.items()}
    res.revision, res.date = str(info.get("RFDS_REVISION", "")), str(info.get("RFDS_DATE", ""))
    b = _Builder(site_id, doc_id, catalog, "XLSX", 0.99, res.revision)
    start = next(i for i, r in enumerate(rows) if r and r[0] == "SECTOR")
    for i, r in enumerate(rows[start + 1:], start + 2):
        if not r or not r[0]:
            continue
        g = lambda c: r[hdr[c]] if c in hdr else None
        kind_hint = {"ANT": "passive", "AAU": "air", "RRU": "radio"}.get(str(g("EQUIP_TYPE")).upper())
        act = str(g("ACTION")).upper()
        page = f"RFDS!R{i}"
        common = dict(rc=g("RC_FT"), az=g("AZ_DEG"), tilt=g("MT_DEG"), bands=g("BANDS"), page=page, kind_hint=kind_hint, feeds=g("FEEDS_ANT_POS"))
        if act in ("KEEP", "RETAIN", "REMOVE"):
            res.existing.append(b.item("RFDS_EXISTING", g("SECTOR"), g("POS"), g("MODEL"), "Existing", **common))
        if act in ("KEEP", "RETAIN", "ADD", "NEW", "SWAP"):
            res.final.append(b.item("RFDS_FINAL", g("SECTOR"), g("POS"), g("MODEL"), ACTION_MAP.get(act, act), **common))
    if "BB_PWR" in wb.sheetnames:
        for r in wb["BB_PWR"].iter_rows(min_row=2, values_only=True):
            if r and str(r[0]).upper() == "BASEBAND":
                res.basebands.append({"model": r[1], "action": ACTION_MAP.get(str(r[2]).upper(), r[2]), "qty": int(r[3] or 1)})
            elif r and r[0] and "DC" in str(r[0]).upper():
                res.dc_load_w = _f(r[3])
    if "REV" in wb.sheetnames:
        res.history = [{"rev": r[0], "date": str(r[1]), "note": r[2]} for r in wb["REV"].iter_rows(min_row=2, values_only=True) if r and r[0]]
    return res


# ------------------------------------------------------------------ PDF export
ROW_RE = re.compile(r"^\s*(Alpha|Beta|Gamma|Delta)\s{2,}", re.I)


def _pdf(path: Path, site_id, doc_id, catalog) -> RfdsResult:
    res = RfdsResult()
    res.template = "PDF"
    pages = pdf_pages_text(path)
    section = None
    for pno, text in enumerate(pages, 1):
        for line in text.splitlines():
            u = line.strip().upper()
            if u.startswith("EXISTING CONFIGURATION"):
                section = "EX"
                continue
            if u.startswith("FINAL CONFIGURATION"):
                section = "FIN"
                continue
            if u.startswith("BASEBAND &") or u.startswith("BASEBAND AND"):
                section = "BB"
                continue
            if u.startswith("REVISION HISTORY"):
                section = "REV"
                continue
            cells = [c.strip() for c in re.split(r"\s{2,}", line.strip()) if c.strip()]
            if section is None or not cells:
                kv = [c.strip() for c in re.split(r"\s{2,}", line.strip()) if c.strip()]
                for k in range(0, len(kv) - 1, 2):
                    res.site[kv[k]] = kv[k + 1]
                continue
            if section in ("EX", "FIN") and ROW_RE.match(line) and len(cells) >= 9:
                res.__dict__.setdefault("_rows", []).append((section, pno, cells))
            elif section == "BB" and cells[0].lower() == "baseband" and len(cells) >= 4:
                res.basebands.append({"model": cells[1], "action": ACTION_MAP.get(cells[2].upper(), cells[2]), "qty": int(cells[3])})
            elif section == "BB" and cells[0].lower().startswith("dc load"):
                res.dc_load_w = _f(cells[-1])
            elif section == "REV" and re.match(r"R\d+$", cells[0]):
                res.history.append({"rev": cells[0], "date": cells[1], "note": " ".join(cells[2:])})
    res.revision = res.site.get("RFDS Revision", "")
    res.date = res.site.get("RFDS Date", "")
    b = _Builder(site_id, doc_id, catalog, "PDF_TEXT", 0.97, res.revision)
    for section, pno, cells in res.__dict__.pop("_rows", []):
        sec, az, pos, action, model, rc, tilt = cells[:7]
        radios, bands = (cells[8], cells[9] if len(cells) > 9 else "") if len(cells) >= 9 else ("", "")
        page = f"p{pno}"
        src = "RFDS_EXISTING" if section == "EX" else "RFDS_FINAL"
        status = "Existing" if section == "EX" else ACTION_MAP.get(action.upper(), action)
        target = res.existing if section == "EX" else res.final
        target.append(b.item(src, sec, pos, model, status, rc=rc, az=az, tilt=tilt, bands=bands, page=page))
        for rmodel, ract in _radios(radios):
            rstatus = "Existing" if section == "EX" else ACTION_MAP.get(ract.upper(), "Retain")
            target.append(b.item(src, sec, pos, rmodel, rstatus, rc=rc, az=az, page=page, kind_hint="radio", feeds=pos))
    if not res.existing and not res.final:
        raise ExtractionError(f"No configuration rows found in {path.name}")
    return res
