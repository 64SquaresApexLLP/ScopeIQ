"""File exports: BOM revision workbook (same columns as the Ericsson REV 0 tool export, plus traceability),
REV 0 -> REV n change list, and the stored redline / RFI / discrepancy package per site."""
from __future__ import annotations

import csv
import json
from dataclasses import asdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from scopeiq.common.logging import get_logger, log_call
from scopeiq.domain import SECTOR_NAMES, BomLine, Discrepancy, Redline, Rfi
from scopeiq.reference.loader import ReferenceData

log = get_logger(__name__)

BOM_COLS = ["Site ID", "Market", "Sector", "Technology", "Band", "Equipment Category", "Manufacturer", "Model", "Manufacturer P/N",
            "Customer P/N", "Description", "UOM", "Design Qty", "Spare Qty", "Total Qty", "Supply (CFM/VFM)", "Action",
            "Source Drawing", "BOM Rule", "Revision", "Status", "Parent", "Reason Code"]
HDR = PatternFill("solid", fgColor="1F3864")


def bom_rows(site_id: str, market: str, lines: list[BomLine], ref: ReferenceData, rev_label: str, status: str) -> list[list]:
    rows = []
    for l in lines:
        it = ref.catalog.get(l.catalog_key)
        par = ref.catalog.get(l.parent_key) if l.parent_key and l.parent_key in ref.catalog else it
        rows.append([site_id, market, SECTOR_NAMES.get(l.sector, "Site"), par.technology or "", par.bands or "", it.category, it.manufacturer,
                     it.model, it.mfr_pn, it.customer_pn, it.description, it.uom, l.design_qty, l.spare_qty, l.total_qty,
                     "N/A (existing)" if l.action == "Remove" else it.supply, l.action, l.source, l.rule_id, rev_label, status,
                     l.parent_key, l.reason_code])
    return rows


@log_call()
def write_bom_xlsx(path: Path, site_id: str, market: str, lines: list[BomLine], ref: ReferenceData, rev_label: str, status: str,
                   changes: list[dict] | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = site_id
    ws.append(BOM_COLS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
    for r in bom_rows(site_id, market, lines, ref, rev_label, status):
        ws.append(r)
    for col, wdt in zip("ABCDEFGHIJKLMNOPQRSTUVW", (10, 7, 8, 9, 9, 18, 14, 22, 18, 12, 40, 6, 9, 9, 9, 10, 8, 30, 10, 8, 12, 10, 16)):
        ws.column_dimensions[col].width = wdt
    if changes is not None:
        ch = wb.create_sheet("Changes vs REV 0")
        ch.append(["Sector", "Part", "Action", "REV 0 Qty", f"{rev_label} Qty", "Delta", "Reason Code(s)", "Discrepancy IDs"])
        for c in ch[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HDR
        for x in changes:
            ch.append([SECTOR_NAMES.get(x["sector"], "Site"), x["part"], x["action"], x["rev0_qty"], x["rev1_qty"], x["delta"],
                       "; ".join(f"{r['reason_code']} ({r['qty']:+g})" for r in x["reasons"]),
                       "; ".join(i for r in x["reasons"] for i in r["disc_ids"])])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    cols = list(rows[0].keys())
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({k: (json.dumps(v, default=str) if isinstance(v, (dict, list)) else v) for k, v in r.items()})


@log_call()
def write_site_package(out_dir: Path, site_id: str, discrepancies: list[Discrepancy], redlines: list[Redline], rfis: list[Rfi]) -> dict:
    """Stored redlines / RFIs / discrepancies for a site as JSON + CSV (also loaded into the database)."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    files = {}
    for name, objs in (("discrepancies", discrepancies), ("redlines", redlines), ("rfis", rfis)):
        rows = [asdict(o) for o in objs]
        (out_dir / f"{site_id}_{name}.json").write_text(json.dumps(rows, indent=2, default=str), encoding="utf-8")
        _csv(out_dir / f"{site_id}_{name}.csv", rows)
        files[name] = str(out_dir / f"{site_id}_{name}.json")
    return files
