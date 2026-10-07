"""Regenerate reference_data/material_catalog.csv and trunk_steps.csv from Data/Reference-Data/*.xlsx and check
that every rule in bom_rules.xlsx exists in reference_data/kit_rules.csv with the same child item.

    python scripts/build_reference_data.py
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scopeiq.config import get_settings  # noqa: E402

COLS = ["CATALOG_KEY", "CATEGORY", "SUBCATEGORY", "MANUFACTURER", "MODEL", "MANUFACTURER_PN", "CUSTOMER_PN", "DESCRIPTION",
        "UOM", "SUPPLY", "CUSTOMER_APPROVED", "UNIT_COST_USD", "RF_PORTS", "POWER_W", "BANDS", "TECHNOLOGY", "HEIGHT_M",
        "WIDTH_M", "DEPTH_M", "ALIASES", "VERSION", "EFFECTIVE_FROM", "EFFECTIVE_TO"]


def aliases(model: str) -> str:
    """Spellings seen on drawings/RFDS for the same model (OCR and A&E variants)."""
    out = {model.upper(), model.upper().replace(" ", ""), model.upper().replace("RADIO ", "")}
    m = re.match(r"AIR (\d+)", model.upper())
    if m:
        out.add(f"AIR{m.group(1)}")
    return "|".join(sorted(a for a in out if a and a != model.upper()))


def main() -> int:
    s = get_settings()
    src = s.path("paths.reference_xlsx_dir")
    dst = s.path("paths.reference_data_dir")
    wb = openpyxl.load_workbook(src / "material_catalog.xlsx", data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0] or str(r[0]).startswith("SYNTHETIC"):
            continue
        vals = ["" if v is None else v for v in r[:19]]
        rows.append(vals + [aliases(str(vals[4])), 1, "2026-01-01", ""])
    with open(dst / "material_catalog.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        w.writerows(rows)
    print(f"material_catalog.csv: {len(rows)} items")

    rb = openpyxl.load_workbook(src / "bom_rules.xlsx", data_only=True)
    steps = [(int(r[0]), r[1]) for r in rb["Trunk Lengths"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], (int, float))]
    with open(dst / "trunk_steps.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["STEP_FT", "CATALOG_KEY", "RULE_TEXT", "VERSION", "EFFECTIVE_FROM", "EFFECTIVE_TO"])
        for ft, key in steps:
            w.writerow([ft, key, "Required length = 1.10 x (vertical + horizontal run); next catalog step at or above", 1, "2026-01-01", ""])
    print(f"trunk_steps.csv: {len(steps)} steps")

    kit = {r["RULE_ID"]: r for r in csv.DictReader(open(dst / "kit_rules.csv", encoding="utf-8"))}
    problems = []
    for r in rb["BOM Rules"].iter_rows(min_row=2, values_only=True):
        if not r[0] or not str(r[0]).startswith("R-"):
            continue
        rid, child = r[0], str(r[2])
        if rid not in kit:
            problems.append(f"{rid} missing from kit_rules.csv")
        elif kit[rid]["CHILD_KEY"] != child:
            problems.append(f"{rid}: xlsx child {child} != kit_rules.csv {kit[rid]['CHILD_KEY']}")
    if problems:
        print("KIT RULE MISMATCH:\n  " + "\n  ".join(problems))
        return 1
    print(f"kit_rules.csv matches bom_rules.xlsx ({len(kit)} rules)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
